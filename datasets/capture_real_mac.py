"""Mac 上截真实应用窗口（数据集采集辅助，不是产品功能）。

和 capture_mac.py 的区别：画面是真实的 Terminal / Numbers 等应用窗口，不是自制 HTML 仿界面。
内容仍然是受控构造的（测试脚本、合成订单、示例表格），来源如实写进 capture_log.csv，source_kind=real_app。

本机限制（2026-10-18 实测）：
- KiroCrew 没有"自动化"权限，不能用 AppleScript 移动窗口、打字、点菜单；
- 代理的沙箱不允许让其他应用打开文档（`open -a Terminal run.command` 会弹出 "Sandbox is not allowed to open
  documents in Terminal"），所以 Terminal / Numbers 窗口要由用户在自己的终端里打开一次；
- 但截图（screencapture -l）和列窗口（Quartz）在沙箱里可用。
因此分两步：
1. `prepare`：生成工作目录、分步脚本 run.command 和示例表格，打印给用户执行的那一条 `open` 命令；
2. 用户执行后，`attach`：按窗口标题找到这些窗口截图。Terminal 的分步脚本每一步真实运行命令后写 ready.N，
   等本工具截完图写 go.N 再继续，于是同一个 Terminal 窗口里的输出逐步变化；窗口大小用终端控制序列
   CSI 8;行;列 t 调整（Terminal 支持，不需要自动化权限）。Numbers 只截打开后的样子，不能点单元格、不能滚动。
窗口用 Quartz 列表按进程名找，截图用 `screencapture -l <窗口号>`；截图后做非空白检查，不覆盖已有文件。

工作目录名、文件名都不带任务编号或场景名（会被截进画面、被 OCR 读到，变成类别线索）。

    data-dataset/.tools/venv/bin/python datasets/capture_real_mac.py terminal prepare --task cao-code-01
    # 用户在自己的终端执行打印出的 open 命令，然后：
    data-dataset/.tools/venv/bin/python datasets/capture_real_mac.py terminal attach --task cao-code-01 \
        --out data-dataset/raw/cao --seq-start 11
"""
from __future__ import annotations

import argparse
import hashlib
import shlex
import shutil
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import Quartz

sys.path.insert(0, str(Path(__file__).resolve().parent))
from capture_mac import append_log, looks_blank  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
PY = REPO / ".venv" / "bin" / "python"
NUMBERS_APP = "Numbers Creator Studio"


def windows(owner: str) -> list[tuple[int, str, tuple[int, int, int, int]]]:
    opts = Quartz.kCGWindowListOptionOnScreenOnly | Quartz.kCGWindowListExcludeDesktopElements
    out = []
    for w in Quartz.CGWindowListCopyWindowInfo(opts, Quartz.kCGNullWindowID) or []:
        if w.get("kCGWindowLayer") == 0 and w.get("kCGWindowOwnerName") == owner:
            b = w["kCGWindowBounds"]
            out.append((int(w["kCGWindowNumber"]), w.get("kCGWindowName", "") or "",
                        (int(b["X"]), int(b["Y"]), int(b["Width"]), int(b["Height"]))))
    return out


def wait_new_window(owner: str, before: set[int], timeout: float = 30.0, min_w: int = 400):
    deadline = time.time() + timeout
    while time.time() < deadline:
        new = [w for w in windows(owner) if w[0] not in before and w[2][2] >= min_w and w[2][3] >= 300]
        if new:
            return max(new, key=lambda w: w[0])
        time.sleep(0.5)
    return None


def window_by_id(owner: str, wid: int):
    return next((w for w in windows(owner) if w[0] == wid), None)


def shoot(task_dir: Path, task_id: str, seq: int, wid: int) -> tuple[Path, str | None]:
    now = datetime.now()
    dest = task_dir / f"{task_id}_{now:%Y%m%d-%H%M%S}_{seq:03d}.png"
    if dest.exists():
        raise SystemExit(f"{dest} 已存在，不覆盖")
    subprocess.run(["screencapture", "-x", "-o", f"-l{wid}", str(dest)], check=True)
    return dest, looks_blank(dest)


SCENES = REPO / "data-dataset" / "scenes-real"


def neutral_dir(task_id: str, scene: str, *, fresh: bool) -> Path:
    d = SCENES / ("w" + hashlib.blake2b(f"{scene}/{task_id}".encode(), digest_size=3).hexdigest())
    if fresh:
        if d.exists():
            shutil.rmtree(d)   # 只清本工具自己生成的工作目录（data-dataset/scenes-real/w*）
        d.mkdir(parents=True)
    elif not d.exists():
        raise SystemExit(f"{d} 不存在：先运行 prepare")
    return d


def window_titled(owner: str, token: str):
    hits = [w for w in windows(owner) if token in w[1]]
    return max(hits, key=lambda w: w[0]) if hits else None


# ---------------------------------------------------------------- Terminal

def terminal_steps() -> list[dict]:
    """cao-code-01 的真实 Terminal 版本：同一个训练脚本改参数重跑。每步：命令、截图前窗口大小（行, 列）、说明。"""
    return [
        dict(cmd="python train.py --batch_size 16 --img 224", size=(36, 120), language="en",
             variation="img=224 → view 形状报错 [16, 3, 224, 224]；窗口 120x36"),
        dict(cmd="python train.py --batch_size 16 --lr 3e-4 --img 200 --epochs 3", size=(36, 120), language="en",
             variation="img 改回 200 重跑，训练成功 3 个 epoch；上一条报错仍在窗口上方"),
        dict(cmd="python train.py --batch_size 16 --img 200 --label_bs 8", size=(36, 120), language="en",
             variation="标签 batch 8 → cross_entropy batch 不一致 ValueError"),
        dict(cmd="python train.py --simulate-oom --batch_size 128", size=(30, 100), language="en",
             variation="batch 128，脚本打印 CUDA OOM 文本（Mac 无 CUDA，不是真实显存溢出）；窗口缩小为 100x30"),
        dict(cmd="python train.py --batch_size 64 --img 256", size=(24, 90), language="en",
             variation="img 256 → 形状报错 [64, 3, 256, 256]；窗口缩小为 90x24"),
    ]


def prepare_terminal(work: Path) -> None:
    shutil.copy(REPO / "datasets" / "scenes" / "code01_train.py", work / "train.py")


def write_command(work: Path, steps: list[dict]) -> Path:
    """分步脚本。提示符 `w… %` 由脚本打印（不是交互 shell），命令本身真实运行，输出原样显示。"""
    venv_bin = shlex.quote(str(PY.parent))
    lines = ["#!/bin/zsh", f"cd {shlex.quote(str(work))}", f"export PATH={venv_bin}:$PATH",
             "export PYTHONDONTWRITEBYTECODE=1", "clear"]
    for i, st in enumerate(steps, 1):
        rows, cols = st["size"]
        lines += [f"printf '\\e[8;{rows};{cols}t'", "sleep 0.5",
                  f"print -r -- {shlex.quote(work.name + ' % ' + st['cmd'])}",
                  st["cmd"],
                  f"print -rn -- {shlex.quote(work.name + ' % ')}",
                  f"touch ready.{i}",
                  f"while [ ! -f go.{i} ]; do sleep 0.2; done",
                  "print"]
    lines += ["print '[采集脚本结束，可以关闭这个窗口]'", "exit 0"]
    p = work / "run.command"
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    p.chmod(0o755)
    return p


def prepare_terminal_scene(task_id: str) -> str:
    work = neutral_dir(task_id, "terminal", fresh=True)
    prepare_terminal(work)
    script = write_command(work, terminal_steps())
    return f"open {shlex.quote(str(script))}"


def run_terminal(task_id: str, out: Path, seq_start: int) -> int:
    task_dir = out / task_id
    task_dir.mkdir(parents=True, exist_ok=True)
    work = neutral_dir(task_id, "terminal", fresh=False)
    steps = terminal_steps()
    win = window_titled("Terminal", work.name) or window_titled("Terminal", "run.command")
    if not win:
        print(f"!! 找不到标题含 {work.name} 或 run.command 的 Terminal 窗口；当前：{windows('Terminal')}", file=sys.stderr)
        return 1
    fails = 0
    for i, st in enumerate(steps, 1):
        if (work / f"go.{i}").exists():
            continue   # 上次已截过这一步
        deadline = time.time() + 120
        while not (work / f"ready.{i}").exists():
            if time.time() > deadline:
                print(f"!! 第 {i} 步超时", file=sys.stderr)
                return fails + 1
            time.sleep(0.3)
        time.sleep(1.0)   # 等 Terminal 重绘
        cur = window_by_id("Terminal", win[0])
        if not cur:
            print("!! Terminal 窗口不见了", file=sys.stderr)
            return fails + 1
        dest, bad = shoot(task_dir, task_id, seq_start + i - 1, cur[0])
        if bad:
            fails += 1
            print(f"!! {dest.name}: {bad}", file=sys.stderr)
        append_log(task_dir, {
            "file": dest.name, "captured_local": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "time_source": "manual", "app": "Terminal", "window_title": cur[1], "language": st["language"],
            "source": f"真实 Terminal 窗口；分步脚本 {work.name}/run.command 真实运行：{st['cmd']}",
            "variation": f"窗口 {cur[2][2]}x{cur[2][3]}pt；{st['variation']}",
            "note": "受控测试内容（测试训练脚本）；提示符由脚本打印，命令与输出为真实运行；窗口由用户用 open 打开"
                    + (f"；{bad}" if bad else ""),
            "source_kind": "real_app"})
        print(f"  ok {dest.name}  [{cur[1][:60]}]")
        (work / f"go.{i}").touch()
    return fails


# ---------------------------------------------------------------- Numbers

BUDGET_CATS = ["Rent 房租", "Groceries 买菜", "Transport 交通", "Phone 话费", "Utilities 水电", "Dining 外食",
               "Books 书", "Gym 健身"]
BUDGET_PLAN = [1850, 450, 140, 40, 110, 280, 50, 89]
JAN = [1850.00, 420.00, 128.50, 38.00, 96.40, 310.00, 64.90, 89.00]


def make_workbooks(work: Path) -> list[tuple[Path, str, str]]:
    """cao-table-01 的真实 Numbers 版本：虚构个人预算表，实际 / 预算 / 差额 / 合计均为公式。需要 openpyxl。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def book(name: str, months: list[tuple[str, list[float]]]) -> Path:
        wb = Workbook()
        wb.remove(wb.active)
        for title, spent in months:
            ws = wb.create_sheet(title)
            ws.append(["Category 类别", "Spent 实际", "Budget 预算", "Diff 差额"])
            for n, (c, v, b) in enumerate(zip(BUDGET_CATS, spent, BUDGET_PLAN), 2):
                ws.append([c, v, b, f"=B{n}-C{n}"])
            last = len(BUDGET_CATS) + 1
            ws.append(["Total 合计", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})", f"=B{last + 1}-C{last + 1}"])
            for c in ws[1] + ws[last + 1]:
                c.font = Font(bold=True)
            for c in ws[1]:
                c.fill = PatternFill("solid", fgColor="DDEBF7")
            for row in ws.iter_rows(min_row=2, min_col=2, max_col=4):
                for c in row:
                    c.number_format = "#,##0.00"
            ws.column_dimensions["A"].width = 20
        p = work / name
        wb.save(p)
        return p

    jan2 = [465.00 if c.startswith("Groceries") else v for c, v in zip(BUDGET_CATS, JAN)]
    feb = [round(v * f, 2) for v, f in zip(JAN, [1, 0.93, 1.1, 1, 1.21, 0.8, 0, 1])]
    return [
        (book("household-2026-jan.xlsx", [("Jan 2026", JAN)]), "mixed", "1 月初始数据（Groceries 420.00）"),
        (book("household-2026-jan-rev.xlsx", [("Jan 2026", jan2)]), "mixed",
         "1 月改金额版本：Groceries 420.00 → 465.00，合计由公式重算"),
        (book("household-2026-feb.xlsx", [("Feb 2026", feb)]), "mixed", "2 月数据（Books 0）"),
    ]


def prepare_numbers_scene(task_id: str, app: str) -> str:
    work = neutral_dir(task_id, "numbers", fresh=True)
    books = make_workbooks(work)
    return f"open -a {shlex.quote(app)} " + " ".join(shlex.quote(str(p)) for p, _, _ in books)


def run_numbers(task_id: str, out: Path, seq_start: int, app: str) -> int:
    task_dir = out / task_id
    task_dir.mkdir(parents=True, exist_ok=True)
    work = neutral_dir(task_id, "numbers", fresh=False)
    books = make_workbooks(work)   # 只为取顺序和说明；文件内容与 prepare 时相同（同一脚本、无随机）
    fails = 0
    for k, (path, lang, desc) in enumerate(books):
        cur = window_titled("Numbers", path.stem)
        if not cur:
            print(f"!! 找不到标题含 {path.stem} 的 Numbers 窗口；当前：{windows('Numbers')}", file=sys.stderr)
            fails += 1
            continue
        dest, bad = shoot(task_dir, task_id, seq_start + k, cur[0])
        if bad:
            fails += 1
        append_log(task_dir, {
            "file": dest.name, "captured_local": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "time_source": "manual", "app": "Numbers", "window_title": cur[1], "language": lang,
            "source": f"真实 Numbers 窗口打开 {work.name}/{path.name}（脚本生成的虚构数据）",
            "variation": f"窗口 {cur[2][2]}x{cur[2][3]}pt；{desc}",
            "note": "虚构数据；未点击、未编辑；文件由用户用 open 打开" + (f"；{bad}" if bad else ""),
            "source_kind": "real_app"})
        print(f"  ok {dest.name}  [{cur[1][:60]}]")
    return fails


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("scene", choices=("terminal", "numbers"))
    ap.add_argument("step", choices=("prepare", "attach"))
    ap.add_argument("--task", required=True)
    ap.add_argument("--out", help="attach 时必填：采集者根目录，如 data-dataset/raw/cao")
    ap.add_argument("--seq-start", type=int, default=1)
    ap.add_argument("--numbers-app", default=NUMBERS_APP)
    a = ap.parse_args()
    if a.step == "prepare":
        cmd = prepare_terminal_scene(a.task) if a.scene == "terminal" else prepare_numbers_scene(a.task, a.numbers_app)
        print("请在你自己的终端里执行：\n" + cmd)
        return 0
    out = Path(a.out).resolve()
    fails = (run_terminal(a.task, out, a.seq_start) if a.scene == "terminal"
             else run_numbers(a.task, out, a.seq_start, a.numbers_app))
    print(f"完成，问题 {fails} 个")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
